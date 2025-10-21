"""Unit tests for export functionality."""

import tempfile
from pathlib import Path
from uuid import uuid4

import pytest

from src.domain.export.exceptions import UnsupportedFormatException
from src.domain.export.manager import ExportManager, create_default_export_manager
from src.domain.export.models import (
    ExportFormat,
    ExportMetrics,
    ExportOptions,
    ExportRequest,
    ExportResult,
    ExportStatus,
    TableOfContents,
)
from src.domain.models.report import DomainReport, ReportSection, ReportType
from src.infrastructure.exporters.html_exporter import HTMLExporter
from src.infrastructure.exporters.markdown_exporter import MarkdownExporter
from src.infrastructure.exporters.pdf_exporter import PDFExporter


@pytest.fixture
def sample_report():
    """Create sample report."""
    report = DomainReport(
        report_id=uuid4(),
        report_title="Test Report on Artificial Intelligence",
        report_type=ReportType.COMPREHENSIVE,
    )

    # Add sections
    section1 = ReportSection(
        section_number="1",
        section_title="Introduction",
        hierarchy_level=0,
        mapped_class_uri="http://example.org/intro",
    )
    section1.add_content({
        "type": "text",
        "text": "This is the introduction section. It provides an overview of artificial intelligence.",
    })
    report.add_section(section1)

    section2 = ReportSection(
        section_number="2",
        section_title="Machine Learning",
        hierarchy_level=0,
        mapped_class_uri="http://example.org/ml",
    )
    section2.add_content({
        "type": "text",
        "text": "Machine learning is a subset of AI. It enables systems to learn from data.",
    })
    report.add_section(section2)

    section2_1 = ReportSection(
        section_number="2.1",
        section_title="Supervised Learning",
        hierarchy_level=1,
        mapped_class_uri="http://example.org/supervised",
    )
    section2_1.add_content({
        "type": "text",
        "text": "Supervised learning uses labeled training data.",
    })
    report.add_section(section2_1)

    # Set metrics
    report.quality_score = 95.5
    report.completeness_score = 98.0

    return report


@pytest.fixture
def export_options():
    """Create export options."""
    return ExportOptions(
        include_toc=True,
        include_metadata=True,
        include_statistics=True,
        theme="light",
    )


@pytest.fixture
def temp_dir():
    """Create temporary directory for test files."""
    with tempfile.TemporaryDirectory() as tmpdir:
        yield Path(tmpdir)


class TestExportModels:
    """Test export domain models."""

    def test_export_options_defaults(self):
        """Test default export options."""
        options = ExportOptions()

        assert options.include_toc is True
        assert options.include_metadata is True
        assert options.theme == "light"
        assert options.font_size == 11

    def test_export_request_creation(self):
        """Test creating export request."""
        request = ExportRequest(
            report_id=uuid4(),
            format=ExportFormat.MARKDOWN,
            output_path=Path("/tmp/test.md"),
        )

        assert request.format == ExportFormat.MARKDOWN
        assert request.output_path == Path("/tmp/test.md")

    def test_export_result_mark_completed(self, temp_dir):
        """Test marking export as completed."""
        result = ExportResult(format=ExportFormat.HTML)

        output_path = temp_dir / "test.html"
        output_path.write_text("test content")

        result.mark_completed(output_path, 100, 1.5)

        assert result.status == ExportStatus.COMPLETED
        assert result.is_successful()
        assert result.file_size == 100
        assert result.export_duration == 1.5

    def test_export_result_mark_failed(self):
        """Test marking export as failed."""
        result = ExportResult(format=ExportFormat.PDF)

        result.mark_failed("Test error")

        assert result.status == ExportStatus.FAILED
        assert not result.is_successful()
        assert result.error_message == "Test error"

    def test_export_result_add_warning(self):
        """Test adding warnings."""
        result = ExportResult()

        result.add_warning("Warning 1")
        result.add_warning("Warning 2")

        assert len(result.warnings) == 2
        assert "Warning 1" in result.warnings

    def test_export_result_file_size_mb(self):
        """Test file size conversion to MB."""
        result = ExportResult()
        result.file_size = 1024 * 1024  # 1 MB

        assert result.get_file_size_mb() == 1.0

    def test_table_of_contents_creation(self):
        """Test creating table of contents."""
        toc = TableOfContents(max_depth=3)

        toc.add_entry("Introduction", 1, anchor="intro")
        toc.add_entry("Background", 2, anchor="background")
        toc.add_entry("Deep Detail", 4, anchor="deep")  # Should be filtered

        assert toc.get_entry_count() == 2  # Level 4 filtered out

    def test_toc_to_markdown(self):
        """Test TOC markdown generation."""
        toc = TableOfContents()

        toc.add_entry("Section 1", 1, anchor="section-1")
        toc.add_entry("Section 1.1", 2, anchor="section-1-1")

        markdown = toc.to_markdown()

        assert "## Table of Contents" in markdown
        assert "Section 1" in markdown
        assert "#section-1" in markdown

    def test_toc_to_html(self):
        """Test TOC HTML generation."""
        toc = TableOfContents()

        toc.add_entry("Section 1", 1, anchor="section-1")
        toc.add_entry("Section 1.1", 2, anchor="section-1-1")

        html = toc.to_html()

        assert '<nav class="table-of-contents">' in html
        assert "<h2>Table of Contents</h2>" in html
        assert 'href="#section-1"' in html


class TestExportMetrics:
    """Test export metrics."""

    def test_metrics_initialization(self):
        """Test metrics initialization."""
        metrics = ExportMetrics()

        assert metrics.total_exports == 0
        assert metrics.successful_exports == 0
        assert metrics.get_success_rate() == 0.0

    def test_record_successful_export(self):
        """Test recording successful export."""
        metrics = ExportMetrics()

        metrics.record_export(
            success=True,
            format=ExportFormat.MARKDOWN,
            duration=1.5,
            file_size=1024,
        )

        assert metrics.total_exports == 1
        assert metrics.successful_exports == 1
        assert metrics.get_success_rate() == 1.0
        assert metrics.average_duration == 1.5
        assert metrics.exports_by_format["markdown"] == 1

    def test_record_failed_export(self):
        """Test recording failed export."""
        metrics = ExportMetrics()

        metrics.record_export(
            success=False,
            format=ExportFormat.PDF,
            duration=0.5,
            file_size=0,
        )

        assert metrics.total_exports == 1
        assert metrics.failed_exports == 1
        assert metrics.get_success_rate() == 0.0

    def test_average_duration_calculation(self):
        """Test average duration calculation."""
        metrics = ExportMetrics()

        metrics.record_export(True, ExportFormat.MARKDOWN, 1.0, 100)
        metrics.record_export(True, ExportFormat.HTML, 2.0, 200)
        metrics.record_export(True, ExportFormat.PDF, 3.0, 300)

        assert metrics.average_duration == 2.0

    def test_total_size_mb(self):
        """Test total size calculation."""
        metrics = ExportMetrics()

        metrics.record_export(True, ExportFormat.MARKDOWN, 1.0, 1024 * 1024)  # 1 MB

        assert metrics.get_total_size_mb() == 1.0


class TestMarkdownExporter:
    """Test Markdown exporter."""

    @pytest.mark.asyncio
    async def test_export_to_markdown(self, sample_report, export_options, temp_dir):
        """Test exporting to Markdown."""
        exporter = MarkdownExporter()
        output_path = temp_dir / "test_report.md"

        result = await exporter.export(sample_report, output_path, export_options)

        assert result.is_successful()
        assert output_path.exists()
        assert result.file_size > 0

        # Check content
        content = output_path.read_text()
        assert "# Test Report on Artificial Intelligence" in content
        assert "## Table of Contents" in content
        assert "## Metadata" in content
        assert "Introduction" in content
        assert "Machine Learning" in content

    @pytest.mark.asyncio
    async def test_markdown_without_toc(self, sample_report, temp_dir):
        """Test Markdown export without TOC."""
        exporter = MarkdownExporter()
        options = ExportOptions(include_toc=False)
        output_path = temp_dir / "no_toc.md"

        result = await exporter.export(sample_report, output_path, options)

        content = output_path.read_text()
        assert "## Table of Contents" not in content

    @pytest.mark.asyncio
    async def test_markdown_auto_extension(self, sample_report, export_options, temp_dir):
        """Test automatic extension addition."""
        exporter = MarkdownExporter()
        output_path = temp_dir / "report"  # No extension

        result = await exporter.export(sample_report, output_path, export_options)

        assert result.output_path.suffix == ".md"

    def test_supports_format(self):
        """Test format support check."""
        exporter = MarkdownExporter()

        assert exporter.supports_format("markdown")
        assert exporter.supports_format("md")
        assert not exporter.supports_format("html")

    def test_get_default_extension(self):
        """Test getting default extension."""
        exporter = MarkdownExporter()

        assert exporter.get_default_extension() == ".md"


class TestHTMLExporter:
    """Test HTML exporter."""

    @pytest.mark.asyncio
    async def test_export_to_html(self, sample_report, export_options, temp_dir):
        """Test exporting to HTML."""
        exporter = HTMLExporter()
        output_path = temp_dir / "test_report.html"

        result = await exporter.export(sample_report, output_path, export_options)

        assert result.is_successful()
        assert output_path.exists()
        assert result.file_size > 0

        # Check content
        content = output_path.read_text()
        assert "<!DOCTYPE html>" in content
        assert "<html" in content
        assert "Test Report on Artificial Intelligence" in content
        assert "table-of-contents" in content
        assert "Introduction" in content

    @pytest.mark.asyncio
    async def test_html_dark_theme(self, sample_report, temp_dir):
        """Test HTML export with dark theme."""
        exporter = HTMLExporter()
        options = ExportOptions(theme="dark")
        output_path = temp_dir / "dark.html"

        result = await exporter.export(sample_report, output_path, options)

        content = output_path.read_text()
        # Dark theme colors should be in CSS
        assert "#1e1e1e" in content or "dark" in content.lower()

    @pytest.mark.asyncio
    async def test_html_escaping(self, temp_dir):
        """Test HTML special character escaping."""
        report = DomainReport(
            report_title="Report with <special> & \"characters\"",
            report_type=ReportType.COMPREHENSIVE,
        )

        section = ReportSection(
            section_number="1",
            section_title="Section <test>",
            hierarchy_level=0,
            mapped_class_uri="test",
        )
        section.add_content({"type": "text", "text": "Content with <tags> & symbols"})
        report.add_section(section)

        exporter = HTMLExporter()
        output_path = temp_dir / "escaped.html"

        result = await exporter.export(report, output_path, ExportOptions())

        content = output_path.read_text()
        assert "&lt;special&gt;" in content
        assert "&amp;" in content
        assert "&quot;" in content

    def test_supports_format(self):
        """Test format support check."""
        exporter = HTMLExporter()

        assert exporter.supports_format("html")
        assert exporter.supports_format("htm")
        assert not exporter.supports_format("pdf")

    def test_get_default_extension(self):
        """Test getting default extension."""
        exporter = HTMLExporter()

        assert exporter.get_default_extension() == ".html"


class TestPDFExporter:
    """Test PDF exporter."""

    @pytest.mark.asyncio
    async def test_export_to_pdf(self, sample_report, export_options, temp_dir):
        """Test exporting to PDF."""
        exporter = PDFExporter()
        output_path = temp_dir / "test_report.pdf"

        result = await exporter.export(sample_report, output_path, export_options)

        assert result.is_successful()
        assert output_path.exists()
        assert result.file_size > 0

        # Check it's a PDF file (starts with %PDF)
        content = output_path.read_bytes()
        assert content.startswith(b"%PDF")

    @pytest.mark.asyncio
    async def test_pdf_auto_extension(self, sample_report, export_options, temp_dir):
        """Test automatic extension addition."""
        exporter = PDFExporter()
        output_path = temp_dir / "report"  # No extension

        result = await exporter.export(sample_report, output_path, export_options)

        assert result.output_path.suffix == ".pdf"

    def test_supports_format(self):
        """Test format support check."""
        exporter = PDFExporter()

        assert exporter.supports_format("pdf")
        assert not exporter.supports_format("html")

    def test_get_default_extension(self):
        """Test getting default extension."""
        exporter = PDFExporter()

        assert exporter.get_default_extension() == ".pdf"


class TestExportManager:
    """Test export manager."""

    def test_manager_initialization(self):
        """Test initializing export manager."""
        manager = ExportManager()

        assert len(manager.get_supported_formats()) == 0

    def test_register_exporter(self):
        """Test registering exporters."""
        manager = ExportManager()

        manager.register_exporter(ExportFormat.MARKDOWN, MarkdownExporter())
        manager.register_exporter(ExportFormat.HTML, HTMLExporter())

        assert len(manager.get_supported_formats()) == 2
        assert manager.is_format_supported(ExportFormat.MARKDOWN)
        assert manager.is_format_supported(ExportFormat.HTML)

    def test_get_exporter(self):
        """Test getting exporter."""
        manager = ExportManager()
        markdown_exporter = MarkdownExporter()

        manager.register_exporter(ExportFormat.MARKDOWN, markdown_exporter)

        retrieved = manager.get_exporter(ExportFormat.MARKDOWN)

        assert retrieved == markdown_exporter

    @pytest.mark.asyncio
    async def test_export_with_manager(self, sample_report, export_options, temp_dir):
        """Test exporting through manager."""
        manager = ExportManager()
        manager.register_exporter(ExportFormat.MARKDOWN, MarkdownExporter())

        output_path = temp_dir / "report.md"

        result = await manager.export(
            report=sample_report,
            format=ExportFormat.MARKDOWN,
            output_path=output_path,
            options=export_options,
        )

        assert result.is_successful()
        assert output_path.exists()

    @pytest.mark.asyncio
    async def test_export_unsupported_format(self, sample_report, temp_dir):
        """Test exporting with unsupported format."""
        manager = ExportManager()

        with pytest.raises(UnsupportedFormatException):
            await manager.export(
                report=sample_report,
                format=ExportFormat.MARKDOWN,  # Not registered
                output_path=temp_dir / "test.md",
            )

    @pytest.mark.asyncio
    async def test_export_multiple_formats(self, sample_report, export_options, temp_dir):
        """Test exporting to multiple formats."""
        manager = ExportManager()
        manager.register_exporter(ExportFormat.MARKDOWN, MarkdownExporter())
        manager.register_exporter(ExportFormat.HTML, HTMLExporter())

        formats = [ExportFormat.MARKDOWN, ExportFormat.HTML]

        results = await manager.export_multiple(
            report=sample_report,
            formats=formats,
            output_dir=temp_dir,
            base_filename="test_report",
            options=export_options,
        )

        assert len(results) == 2
        assert ExportFormat.MARKDOWN in results
        assert ExportFormat.HTML in results
        assert all(r.is_successful() for r in results.values())

        # Check files exist
        assert (temp_dir / "test_report.md").exists()
        assert (temp_dir / "test_report.html").exists()

    @pytest.mark.asyncio
    async def test_metrics_tracking(self, sample_report, export_options, temp_dir):
        """Test metrics tracking."""
        manager = ExportManager()
        manager.register_exporter(ExportFormat.MARKDOWN, MarkdownExporter())

        # Perform export
        await manager.export(
            report=sample_report,
            format=ExportFormat.MARKDOWN,
            output_path=temp_dir / "test.md",
            options=export_options,
        )

        metrics = manager.get_metrics()

        assert metrics.total_exports == 1
        assert metrics.successful_exports == 1
        assert metrics.exports_by_format["markdown"] == 1

    def test_get_metrics_summary(self):
        """Test getting metrics summary."""
        manager = ExportManager()

        summary = manager.get_metrics_summary()

        assert "total_exports" in summary
        assert "success_rate" in summary
        assert "average_duration" in summary

    def test_create_default_export_manager(self):
        """Test creating manager with default exporters."""
        manager = create_default_export_manager()

        assert manager.is_format_supported(ExportFormat.MARKDOWN)
        assert manager.is_format_supported(ExportFormat.HTML)
        assert manager.is_format_supported(ExportFormat.PDF)


class TestExportEdgeCases:
    """Test edge cases and error handling."""

    @pytest.mark.asyncio
    async def test_export_empty_report(self, export_options, temp_dir):
        """Test exporting empty report."""
        report = DomainReport(
            report_title="Empty Report",
            report_type=ReportType.COMPREHENSIVE,
        )

        exporter = MarkdownExporter()
        output_path = temp_dir / "empty.md"

        result = await exporter.export(report, output_path, export_options)

        assert result.is_successful()
        assert output_path.exists()

    @pytest.mark.asyncio
    async def test_export_with_special_characters_in_filename(
        self, sample_report, export_options, temp_dir
    ):
        """Test export with special characters in filename."""
        exporter = MarkdownExporter()
        output_path = temp_dir / "report with spaces & special.md"

        result = await exporter.export(sample_report, output_path, export_options)

        assert result.is_successful()
        assert output_path.exists()


class TestDOCXExporter:
    """Test DOCX exporter."""

    @pytest.mark.asyncio
    async def test_export_to_docx(self, sample_report, export_options, temp_dir):
        """Test exporting to DOCX."""
        from src.infrastructure.exporters.docx_exporter import DOCXExporter

        exporter = DOCXExporter()
        output_path = temp_dir / "test_report.docx"

        result = await exporter.export(sample_report, output_path, export_options)

        assert result.is_successful()
        assert output_path.exists()
        assert result.file_size > 0

        # Check it's a DOCX file (ZIP format)
        import zipfile
        assert zipfile.is_zipfile(output_path)

    @pytest.mark.asyncio
    async def test_docx_auto_extension(self, sample_report, export_options, temp_dir):
        """Test automatic extension addition."""
        from src.infrastructure.exporters.docx_exporter import DOCXExporter

        exporter = DOCXExporter()
        output_path = temp_dir / "report"  # No extension

        result = await exporter.export(sample_report, output_path, export_options)

        assert result.output_path.suffix == ".docx"

    @pytest.mark.asyncio
    async def test_docx_without_toc(self, sample_report, temp_dir):
        """Test DOCX export without TOC."""
        from src.infrastructure.exporters.docx_exporter import DOCXExporter

        exporter = DOCXExporter()
        options = ExportOptions(include_toc=False)
        output_path = temp_dir / "no_toc.docx"

        result = await exporter.export(sample_report, output_path, options)

        assert result.is_successful()
        assert output_path.exists()

    @pytest.mark.asyncio
    async def test_docx_with_statistics(self, sample_report, temp_dir):
        """Test DOCX export with statistics."""
        from src.infrastructure.exporters.docx_exporter import DOCXExporter

        exporter = DOCXExporter()
        options = ExportOptions(include_statistics=True)
        output_path = temp_dir / "with_stats.docx"

        result = await exporter.export(sample_report, output_path, options)

        assert result.is_successful()

    def test_docx_supports_format(self):
        """Test format support check."""
        from src.infrastructure.exporters.docx_exporter import DOCXExporter

        exporter = DOCXExporter()

        assert exporter.supports_format("docx")
        assert exporter.supports_format("doc")
        assert not exporter.supports_format("pdf")

    def test_docx_get_default_extension(self):
        """Test getting default extension."""
        from src.infrastructure.exporters.docx_exporter import DOCXExporter

        exporter = DOCXExporter()

        assert exporter.get_default_extension() == ".docx"


class TestXLSXExporter:
    """Test XLSX exporter."""

    @pytest.mark.asyncio
    async def test_export_to_xlsx(self, sample_report, export_options, temp_dir):
        """Test exporting to XLSX."""
        from src.infrastructure.exporters.xlsx_exporter import XLSXExporter

        exporter = XLSXExporter()
        output_path = temp_dir / "test_report.xlsx"

        result = await exporter.export(sample_report, output_path, export_options)

        assert result.is_successful()
        assert output_path.exists()
        assert result.file_size > 0

        # Check it's an XLSX file (ZIP format)
        import zipfile
        assert zipfile.is_zipfile(output_path)

    @pytest.mark.asyncio
    async def test_xlsx_auto_extension(self, sample_report, export_options, temp_dir):
        """Test automatic extension addition."""
        from src.infrastructure.exporters.xlsx_exporter import XLSXExporter

        exporter = XLSXExporter()
        output_path = temp_dir / "report"  # No extension

        result = await exporter.export(sample_report, output_path, export_options)

        assert result.output_path.suffix == ".xlsx"

    @pytest.mark.asyncio
    async def test_xlsx_without_metadata(self, sample_report, temp_dir):
        """Test XLSX export without metadata."""
        from src.infrastructure.exporters.xlsx_exporter import XLSXExporter

        exporter = XLSXExporter()
        options = ExportOptions(include_metadata=False)
        output_path = temp_dir / "no_metadata.xlsx"

        result = await exporter.export(sample_report, output_path, options)

        assert result.is_successful()
        assert output_path.exists()

    @pytest.mark.asyncio
    async def test_xlsx_with_statistics(self, sample_report, temp_dir):
        """Test XLSX export with statistics."""
        from src.infrastructure.exporters.xlsx_exporter import XLSXExporter

        exporter = XLSXExporter()
        options = ExportOptions(include_statistics=True)
        output_path = temp_dir / "with_stats.xlsx"

        result = await exporter.export(sample_report, output_path, options)

        assert result.is_successful()

    def test_xlsx_supports_format(self):
        """Test format support check."""
        from src.infrastructure.exporters.xlsx_exporter import XLSXExporter

        exporter = XLSXExporter()

        assert exporter.supports_format("xlsx")
        assert exporter.supports_format("xls")
        assert not exporter.supports_format("docx")

    def test_xlsx_get_default_extension(self):
        """Test getting default extension."""
        from src.infrastructure.exporters.xlsx_exporter import XLSXExporter

        exporter = XLSXExporter()

        assert exporter.get_default_extension() == ".xlsx"


class TestMSOfficeIntegration:
    """Test MS Office export integration."""

    @pytest.mark.asyncio
    async def test_export_manager_with_office_formats(self, sample_report, export_options, temp_dir):
        """Test export manager with MS Office formats."""
        from src.domain.export.manager import create_default_export_manager

        manager = create_default_export_manager()

        # Export to DOCX
        docx_result = await manager.export(
            report=sample_report,
            format=ExportFormat.DOCX,
            output_path=temp_dir / "report.docx",
            options=export_options,
        )

        assert docx_result.is_successful()
        assert (temp_dir / "report.docx").exists()

    @pytest.mark.asyncio
    async def test_batch_export_with_office(self, sample_report, export_options, temp_dir):
        """Test batch export including MS Office formats."""
        from src.domain.export.manager import create_default_export_manager

        manager = create_default_export_manager()

        formats = [
            ExportFormat.MARKDOWN,
            ExportFormat.HTML,
            ExportFormat.PDF,
            ExportFormat.DOCX,
        ]

        results = await manager.export_multiple(
            report=sample_report,
            formats=formats,
            output_dir=temp_dir,
            base_filename="full_report",
            options=export_options,
        )

        # Check all exports succeeded
        assert len(results) == len(formats)
        for format in formats:
            assert format in results
            assert results[format].is_successful()

        # Check all files exist
        assert (temp_dir / "full_report.md").exists()
        assert (temp_dir / "full_report.html").exists()
        assert (temp_dir / "full_report.pdf").exists()
        assert (temp_dir / "full_report.docx").exists()

    @pytest.mark.asyncio
    async def test_docx_with_complex_content(self, temp_dir):
        """Test DOCX export with complex content."""
        from src.infrastructure.exporters.docx_exporter import DOCXExporter

        # Create report with various content types
        report = DomainReport(
            report_title="Complex Report",
            report_type=ReportType.COMPREHENSIVE,
        )

        section = ReportSection(
            section_number="1",
            section_title="Test Section",
            hierarchy_level=0,
            mapped_class_uri="test",
        )

        # Add different content types
        section.add_content({
            "type": "text",
            "text": "Regular paragraph content.",
        })
        section.add_content({
            "type": "code",
            "text": "def hello():\n    print('Hello')",
            "language": "python",
        })
        section.add_content({
            "type": "quote",
            "text": "This is a quote.",
        })
        section.add_content({
            "type": "table",
            "data": [
                ["Header 1", "Header 2"],
                ["Data 1", "Data 2"],
                ["Data 3", "Data 4"],
            ],
        })

        report.add_section(section)

        exporter = DOCXExporter()
        output_path = temp_dir / "complex.docx"

        result = await exporter.export(report, output_path, ExportOptions())

        assert result.is_successful()
        assert output_path.exists()

    @pytest.mark.asyncio
    async def test_xlsx_multiple_sheets(self, sample_report, export_options, temp_dir):
        """Test XLSX creates multiple worksheets."""
        from src.infrastructure.exporters.xlsx_exporter import XLSXExporter

        exporter = XLSXExporter()
        output_path = temp_dir / "multi_sheet.xlsx"

        result = await exporter.export(sample_report, output_path, export_options)

        assert result.is_successful()

        # Try to load and check sheets
        try:
            from openpyxl import load_workbook
            wb = load_workbook(output_path)

            # Should have multiple sheets
            assert len(wb.sheetnames) >= 3

            # Check expected sheets
            assert "Report Overview" in wb.sheetnames
            assert "Sections" in wb.sheetnames

            if export_options.include_toc:
                assert "Table of Contents" in wb.sheetnames

            if export_options.include_metadata:
                assert "Metadata" in wb.sheetnames

            if export_options.include_statistics:
                assert "Statistics" in wb.sheetnames

        except ImportError:
            # openpyxl not available, skip validation
            pass
